import csv
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import torch

from cf.data import CFData, file_hash
from cf.evaluate import evaluate
from cf.models import build_model
from eval_fism import evaluate_checkpoint, export_excel, main, report_table, select_checkpoints
from search_cf import atomic_json, digest, read_json
from train_cf import parse_args


class FinalFISMTests(unittest.TestCase):
    def setUp(self):
        torch.set_num_threads(2)
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.search = self.root / 'search'
        self.data_root = self.root / 'data'
        directory = self.data_root / 'amazon_baby' / 'cf_data'
        directory.mkdir(parents=True)
        for split, content in dict(train='0 0 1 0\n1 1 2 4 5\n', valid='0 2\n1 3\n',
                                   test='0 3\n1 0\n').items():
            (directory / (split + '.txt')).write_text(content, encoding='utf-8')
        self.data = CFData(directory, 'count')
        self.jobs = []
        self.paths = []
        self.models = {}
        for number, (variant, recall) in enumerate([('original', .3), ('item-only', .4), ('item-only', .8)]):
            args = parse_args(['--dataset', 'Baby', '--model', 'fism', '--variant', variant,
                               '--data-root', str(self.data_root), '--duplicate-items', 'count',
                               '--dim', '4', '--b-mode', 'degree', '--b-a', str(number / 2),
                               '--fism-user-bias', str(number % 2), '--validation-only'])
            config = {key: str(value) if isinstance(value, Path) else value for key, value in vars(args).items()}
            job = dict(id='job' + str(number), dataset='amazon_baby', model='fism',
                       variant=variant, seed=2025, B='B' + str(number), params={})
            model = build_model(self.data, args, 'fism', variant, 'cpu')
            with torch.no_grad():
                model.item_bias.copy_(torch.arange(self.data.n_items, dtype=torch.float32))
            self.models[job['id']] = model
            valid = {'recall@20': recall, 'ndcg@20': .2, 'hr@20': .5, 'users': 2}
            run = self.search / job['dataset'] / job['id'] / 'attempt_fixture' / 'run'
            run.mkdir(parents=True)
            metadata = dict(config=config, data=self.data.stats, smoke_test=False,
                            protocol='AlphaFree full catalog: valid masks train+test; test masks train+valid; target retained')
            torch.save(dict(state_dict=model.state_dict(), epoch=5, metadata=metadata, valid=valid), run / 'best.pt')
            result = dict(dataset=job['dataset'], model='fism', variant=variant, seed=2025,
                          validation_only=True, test=None, smoke_test=False, best_epoch=5,
                          valid=valid, run_dir='/old/machine/unavailable/path')
            atomic_json(run / 'result.json', result)
            atomic_json(run.parent.parent / 'status.json', dict(state='completed',
                        job_digest=digest(job), attempt='attempt_fixture'))
            self.jobs.append(job)
            self.paths.append(run)
        atomic_json(self.search / 'manifest.json', dict(jobs=self.jobs))

    def tearDown(self):
        self.temp.cleanup()

    def test_selection_uses_validation_and_resolves_moved_paths(self):
        chosen = select_checkpoints(self.search)
        self.assertEqual([entry['job']['id'] for entry in chosen], ['job0', 'job2'])
        self.assertTrue(all(entry['checkpoint'].is_file() for entry in chosen))

    def test_incomplete_search_refused(self):
        (self.paths[1] / 'best.pt').unlink()
        with self.assertRaisesRegex(ValueError, 'incomplete'):
            select_checkpoints(self.search)

    def test_smoke_result_not_selected(self):
        path = self.paths[2] / 'result.json'
        result = read_json(path)
        result['smoke_test'] = True
        atomic_json(path, result)
        with self.assertRaisesRegex(ValueError, 'smoke-only'):
            select_checkpoints(self.search)

    def test_reload_matches_direct_test_metrics_without_optimizer(self):
        before = [file_hash(path / 'best.pt') for path in self.paths]
        with patch('torch.optim.Adam', side_effect=AssertionError('Must not train')):
            for entry in select_checkpoints(self.search):
                original = self.models[entry['job']['id']]
                def checked_evaluate(restored, data, split, k, batch_size, max_users):
                    self.assertEqual(split, 'test')
                    self.assertEqual(max_users, 0)
                    torch.testing.assert_close(restored.score_all([0, 1]), original.score_all([0, 1]))
                    return evaluate(restored, data, split, k, batch_size, max_users)
                with patch('eval_fism.evaluate', side_effect=checked_evaluate):
                    result = evaluate_checkpoint(entry, 'cpu', self.data_root, (1, 2), 1)
                expected = {}
                for k in (1, 2):
                    expected.update(evaluate(self.models[entry['job']['id']], self.data, 'test', k, 1))
                self.assertEqual(result['test'], expected)
        self.assertEqual(before, [file_hash(path / 'best.pt') for path in self.paths])

    def test_changed_data_refused_before_evaluation(self):
        selected = select_checkpoints(self.search)[0]
        path = self.data.directory / 'test.txt'
        path.write_text('0 2\n1 0\n', encoding='utf-8')
        with patch('eval_fism.evaluate', side_effect=AssertionError('Must check hashes first')):
            with self.assertRaisesRegex(ValueError, 'hashes differ'):
                evaluate_checkpoint(selected, 'cpu', self.data_root)

    def test_cli_outputs_and_dry_run(self):
        argv = ['--search-output', str(self.search), '--data-root', str(self.data_root),
                '--device', 'cpu', '--datasets', 'Baby', '--top-k', '1', '2']
        with patch('eval_fism.evaluate', side_effect=AssertionError('Dry run must not test')):
            main(argv + ['--dry-run'])
        self.assertFalse((self.search / 'final_test.json').exists())
        results = main(argv)
        self.assertEqual(len(results), 2)
        self.assertEqual(read_json(self.search / 'final_test.json'), results)
        with (self.search / 'final_test.csv').open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 2)
        self.assertIn('test_recall@2', rows[0])
        self.assertTrue(all(read_json(path / 'result.json')['test'] is None for path in self.paths))

    def test_all_models_includes_nonwinning_configurations(self):
        chosen = select_checkpoints(self.search, all_models=True)
        self.assertEqual([entry['job']['id'] for entry in chosen], ['job0', 'job1', 'job2'])
        self.assertEqual([entry['validation_selected'] for entry in chosen], [True, False, True])
        with patch('eval_fism.export_excel') as export:
            results = main(['--search-output', str(self.search), '--all-models', '--device', 'cpu'])
        self.assertEqual(len(results), 3)
        self.assertTrue(all(row['status'] == 'completed' for row in results))
        self.assertEqual(export.call_args[0][0], (self.search / 'all_test.xlsx').resolve())
        self.assertEqual(len(export.call_args[0][1]['rows']), 3)
        self.assertEqual(len(read_json(self.search / 'all_test.json')), 3)

    def test_failure_is_recorded_and_later_models_still_tested(self):
        def fail_one(entry, *args):
            if entry['job']['id'] == 'job1':
                raise RuntimeError('fixture failure')
            return evaluate_checkpoint(entry, *args)
        with patch('eval_fism.evaluate_checkpoint', side_effect=fail_one), patch('eval_fism.export_excel'):
            with self.assertRaisesRegex(RuntimeError, '1 checkpoints failed'):
                main(['--search-output', str(self.search), '--all-models', '--device', 'cpu'])
        results = read_json(self.search / 'all_test.json')
        self.assertEqual([row['status'] for row in results], ['completed', 'failed', 'completed'])
        table = report_table(results, [20])
        self.assertIsNone(table['rows'][1][table['columns'].index('test_recall@20')])
        self.assertIn('fixture failure', table['rows'][1][table['columns'].index('error')])

    def test_export_only_does_not_reload_or_evaluate_checkpoints(self):
        with patch('eval_fism.export_excel'):
            main(['--search-output', str(self.search), '--all-models', '--device', 'cpu', '--top-k', '1', '2'])
        with patch('eval_fism.select_checkpoints', side_effect=AssertionError('No reload')):
            with patch('eval_fism.evaluate_checkpoint', side_effect=AssertionError('No evaluation')):
                with patch('eval_fism.export_excel') as export:
                    main(['--search-output', str(self.search), '--all-models', '--export-only'])
        self.assertIn('test_recall@2', export.call_args[0][1]['columns'])
        self.assertNotIn('test_recall@20', export.call_args[0][1]['columns'])

    def test_excel_retains_all_rows_numeric_metrics_and_filters(self):
        from openpyxl import load_workbook
        results = [evaluate_checkpoint(entry, 'cpu', self.data_root)
                   for entry in select_checkpoints(self.search, all_models=True)]
        table = report_table(results, [20])
        path = self.root / 'all_models.xlsx'
        export_excel(path, table)
        workbook = load_workbook(path, data_only=True)
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet.max_column, len(table['columns']))
        self.assertEqual(sheet.cell(1, 1).value, 'dataset')
        self.assertEqual(sheet.cell(3, table['columns'].index('B') + 1).value, 'B1')
        self.assertIs(sheet.cell(3, table['columns'].index('validation_selected') + 1).value, False)
        self.assertIsInstance(sheet.cell(2, table['columns'].index('test_recall@20') + 1).value, (float, int))
        with ZipFile(path) as archive:
            self.assertEqual(len([name for name in archive.namelist()
                                  if name.startswith('xl/tables/') and name.endswith('.xml')]), 1)
        self.assertTrue(sheet.freeze_panes)
        workbook.close()

    def test_excel_on_server_without_node_runtime(self):
        from openpyxl import load_workbook
        table = dict(columns=['dataset', 'test_recall@20', 'status'],
                     rows=[['baby', .1234567, 'completed'], ['steam', None, 'failed']])
        path = self.root / 'server.xlsx'
        with patch.dict(os.environ, {'FISM_NODE': ''}):
            with patch('eval_fism.shutil.which', return_value=None):
                export_excel(path, table)
        book = load_workbook(path, data_only=True)
        self.assertEqual(book.active['B2'].value, .1234567)
        self.assertIsNone(book.active['B3'].value)
        self.assertEqual(book.active.freeze_panes, 'D2')
        book.close()


if __name__ == '__main__':
    unittest.main()
