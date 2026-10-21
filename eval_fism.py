"""Test saved FISM checkpoints (validation-selected or all configurations)."""
import argparse
import csv
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from types import SimpleNamespace

import torch

from cf.data import CFData, DATASETS, file_hash
from cf.evaluate import evaluate
from cf.models import build_model
from search_cf import atomic_json, completed_result, job_directory, read_json, summarize
from train_cf import ROOT, load_checkpoint


def select_checkpoints(root, datasets=None, variants=None, all_models=False):
    manifest = read_json(root / 'manifest.json')
    jobs = [job for job in manifest['jobs'] if job['model'] == 'fism'
            and (not datasets or job['dataset'] in datasets)
            and (not variants or job['variant'] in variants)]
    if not jobs:
        raise ValueError('No matching FISM jobs in manifest.json')
    if datasets and set(datasets) - {job['dataset'] for job in jobs}:
        raise ValueError('Some requested datasets have no matching FISM jobs')
    if variants and set(variants) - {job['variant'] for job in jobs}:
        raise ValueError('Some requested variants have no matching FISM jobs')
    groups = summarize(root, jobs)
    selected = []
    for group in groups:
        if not group['complete'] or group['best'] is None:
            raise ValueError('Search incomplete or smoke-only: {} / {} / seed {}; '
                             'completed {}/{}. Finish the search before final testing.'.format(
                                 group['dataset'], group['variant'], group['seed'],
                                 group['completed'], group['planned']))
    winners = {group['best']['job']['id'] for group in groups}
    choices = ([dict(job=job, result=completed_result(root, job)) for job in jobs]
               if all_models else [group['best'] for group in groups])
    for winner in choices:
        job, result = winner['job'], winner['result']
        directory = job_directory(root, job)
        status = read_json(directory / 'status.json')
        # Resolve relative to the supplied search folder, not old machine paths in result.json.
        checkpoint = directory / status['attempt'] / 'run' / 'best.pt'
        if not checkpoint.is_file() and result.get('checkpoint_retention') == 'delete_after_test':
            raise ValueError('Checkpoint was deleted after testing. Read saved test metrics in '
                             'result.json / training.txt; this run cannot be re-evaluated: ' + str(checkpoint.parent))
        selected.append(dict(job=job, result=result, checkpoint=checkpoint,
                             validation_selected=job['id'] in winners,
                             selection='all configurations; checkpoint epoch selected by validation Recall@20'
                             if all_models else 'validation Recall@20 within dataset/variant/seed'))
    return selected


def evaluate_checkpoint(selected, device, data_root=None, top_ks=(20,), batch_size=128):
    path = selected['checkpoint']
    checkpoint = load_checkpoint(path, 'cpu')
    metadata = checkpoint['metadata']
    config = metadata['config']
    job = selected['job']
    if metadata.get('smoke_test') or config.get('max_steps') or config.get('eval_users'):
        raise ValueError('Refusing a smoke-test checkpoint: ' + str(path))
    for key in ('model', 'dataset', 'variant', 'seed'):
        if config[key] != job[key]:
            raise ValueError('Checkpoint/job mismatch: ' + key)
    if checkpoint['epoch'] != selected['result']['best_epoch']:
        raise ValueError('Checkpoint epoch differs from saved search result')
    if checkpoint['valid'] != selected['result']['valid']:
        raise ValueError('Checkpoint validation metrics differ from saved search result')
    base = Path(data_root) if data_root is not None else Path(config['data_root'])
    data = CFData(base / config['dataset'] / 'cf_data', config.get('duplicate_items', 'error'))
    if data.stats['sha256'] != metadata['data']['sha256']:
        raise ValueError('Data split hashes differ from training. Restore the original train/valid/test files.')
    args = SimpleNamespace(**config)
    model = build_model(data, args, 'fism', config['variant'], device)
    model.load_state_dict(checkpoint['state_dict'], strict=True)
    metrics = {}
    for k in top_ks:
        metrics.update(evaluate(model, data, 'test', k, batch_size, 0))
    export_config = {key: config.get(key) for key in
                     ('b_mode', 'b_a', 'b_b', 'fism_alpha', 'fism_reg', 'fism_negatives',
                      'fism_user_bias', 'dim', 'lr', 'batch_size', 'duplicate_items')}
    if config['variant'] == 'original':
        export_config.update(b_mode=None, b_a=None, b_b=None)
    else:
        export_config['fism_alpha'] = None  # Not used by FISM-B.
    return dict(dataset=job['dataset'], model='fism', variant=job['variant'], seed=job['seed'],
                B=job['B'], job_id=job['id'], params=job['params'],
                config=export_config, status='completed', validation_selected=selected.get('validation_selected'),
                checkpoint=str(path.resolve()), checkpoint_sha256=file_hash(path),
                best_epoch=checkpoint['epoch'], valid=checkpoint['valid'], test=metrics,
                selection=selected.get('selection', 'validation Recall@20 within dataset/variant/seed'),
                protocol=metadata['protocol'], data_sha256=data.stats['sha256'],
                training_adapter_sha256=metadata.get('adapter_sha256', {}),
                evaluation_source_sha256={name: file_hash(ROOT / name) for name in
                                         ('eval_fism.py', 'cf/fism.py', 'cf/data.py', 'cf/evaluate.py')})


def report_table(results, top_ks):
    fields = ['dataset', 'variant', 'B', 'seed', 'best_epoch', 'status']
    fields += ['test_{}@{}'.format(metric, k) for k in top_ks for metric in ('recall', 'ndcg', 'hr')]
    fields += ['valid_recall@20', 'valid_ndcg@20', 'valid_hr@20', 'validation_selected', 'test_users']
    fields += ['b_mode', 'b_a', 'b_b', 'fism_alpha', 'fism_reg', 'fism_negatives', 'fism_user_bias',
               'dim', 'lr', 'batch_size', 'duplicate_items', 'job_id', 'checkpoint', 'error']
    rows = []
    for result in results:
        row = {key: result.get(key) for key in fields}
        row.update(result.get('config', {}))
        for metric, value in result.get('test', {}).items():
            row['test_users' if metric == 'users' else 'test_' + metric] = value
        for metric, value in result.get('valid', {}).items():
            row['valid_' + metric] = value
        rows.append([row.get(key) for key in fields])
    return dict(columns=fields, rows=rows)


def export_excel(output, table):
    """Prefer Artifact Tool; servers without that runtime can use openpyxl."""
    node = os.environ.get('FISM_NODE') or shutil.which('node')
    builder = ROOT / 'export_fism_xlsx.mjs'
    if node and builder.is_file():
        with tempfile.TemporaryDirectory() as temporary:
            payload = Path(temporary) / 'table.json'
            atomic_json(payload, table)
            process = subprocess.run([node, str(builder), str(payload), str(output)],
                                     stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                     universal_newlines=True, encoding='utf-8', errors='replace')
        if process.returncode == 0:
            return
        if process.returncode != 42:
            raise RuntimeError('Excel export failed: ' + process.stderr)
    # The training server need not have the Codex/Node artifact runtime installed.
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as error:
        raise RuntimeError('Excel export needs openpyxl on this server: python -m pip install openpyxl. '
                           'JSON/CSV are saved; rerun with --export-only after installation.') from error
    book = Workbook()
    sheet = book.active
    sheet.title = 'All model tests'
    sheet.sheet_view.showGridLines = False
    sheet.append(table['columns'])
    for row in table['rows']:
        sheet.append(row)
    for col, field in enumerate(table['columns'], 1):
        letter = get_column_letter(col)
        width = 55 if field in ('checkpoint', 'error') else 42 if field == 'job_id' else 22
        sheet.column_dimensions[letter].width = width
        header = sheet.cell(1, col)
        header.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        header.fill = PatternFill('solid', fgColor='243B53')
        header.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, col)
            cell.font = Font(name='Arial', size=10, color='243B53')
            cell.alignment = Alignment(vertical='center')
            if isinstance(cell.value, str):
                cell.data_type = 's'
            if '@' in field or field in ('fism_reg', 'lr'):
                cell.number_format = '0.000000'
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = 'D2'
    tab = Table(displayName='FISMTests', ref=sheet.dimensions)
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(tab)
    book.save(output)


def save_reports(output, results, top_ks, excel=False):
    atomic_json(output, results)
    table = report_table(results, top_ks)
    with output.with_suffix('.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(table['columns'])
        writer.writerows(table['rows'])
    if excel:
        export_excel(output.with_suffix('.xlsx'), table)
    print('Saved {} / CSV{}'.format(output, ' / Excel' if excel else ''), flush=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--search-output', type=Path, required=True, help='Search folder containing manifest.json')
    parser.add_argument('--data-root', type=Path, help='Override training data root when moving machines')
    parser.add_argument('--datasets', nargs='+', help='Optional aliases, e.g. Baby Steam')
    parser.add_argument('--variants', nargs='+', choices=['original', 'item-only'])
    parser.add_argument('--all-models', action='store_true', help='Test every configuration; also export Excel')
    parser.add_argument('--export-only', action='store_true', help='Rebuild CSV/Excel from saved JSON without testing')
    parser.add_argument('--device', default='auto')
    parser.add_argument('--top-k', type=int, nargs='+', default=[20])
    parser.add_argument('--eval-batch-size', type=int, default=128)
    parser.add_argument('--threads', type=int, default=2)
    parser.add_argument('--output', type=Path, help='JSON path; default: final_test.json or all_test.json')
    parser.add_argument('--dry-run', action='store_true', help='Show checkpoints without testing')
    options = parser.parse_args(argv)
    if min(options.top_k + [options.eval_batch_size, options.threads]) < 1:
        parser.error('top-k, eval-batch-size and threads must be positive')
    if options.output and options.output.suffix.lower() != '.json':
        parser.error('--output must end in .json')
    root = options.search_output.resolve()
    output = options.output or root / ('all_test.json' if options.all_models else 'final_test.json')
    if options.export_only:
        results = read_json(output)
        if not results:
            parser.error('No saved results to export')
        top_ks = sorted({int(key.split('@')[1]) for row in results for key in row.get('test', {}) if '@' in key})
        save_reports(output, results, top_ks or options.top_k, excel=True)
        return results
    datasets = [DATASETS.get(name.lower(), name) for name in options.datasets] if options.datasets else None
    selected = select_checkpoints(root, datasets, options.variants, options.all_models)
    print('Checkpoints to test: {}'.format(len(selected)), flush=True)
    for entry in selected:
        job = entry['job']
        print('SELECTED {} / {} / seed {} / B={} / valid Recall@20={:.6f} / {}'.format(
            job['dataset'], job['variant'], job['seed'], job['B'],
            entry['result']['valid']['recall@20'], entry['checkpoint']), flush=True)
    if options.dry_run:
        return selected
    device = torch.device(('cuda:0' if torch.cuda.is_available() else 'cpu')
                          if options.device == 'auto' else options.device)
    if device.type == 'cuda' and not torch.cuda.is_available():
        raise RuntimeError('CUDA unavailable; use a CUDA Python environment or --device cpu')
    torch.set_num_threads(options.threads)
    results = []
    top_ks = sorted(set(options.top_k))
    for index, entry in enumerate(selected, 1):
        print('EVALUATING {}/{} {}'.format(index, len(selected), entry['job']['id']), flush=True)
        try:
            result = evaluate_checkpoint(entry, device, options.data_root, top_ks, options.eval_batch_size)
        except Exception as error:
            if not options.all_models:
                raise
            job = entry['job']
            result = dict(dataset=job['dataset'], variant=job['variant'], seed=job['seed'], B=job['B'],
                          job_id=job['id'], params=job['params'], status='failed',
                          best_epoch=entry['result']['best_epoch'], valid=entry['result']['valid'], test={},
                          validation_selected=entry['validation_selected'], checkpoint=str(entry['checkpoint']),
                          error='{}: {}'.format(type(error).__name__, error))
            print('FAILED ' + result['error'], flush=True)
        finally:
            if device.type == 'cuda':
                torch.cuda.empty_cache()
        results.append(result)
        atomic_json(output, results)  # Preserve completed evaluations if interrupted or export fails.
        print('TEST ' + json.dumps({key: result[key] for key in
                                   ('dataset', 'variant', 'seed', 'B', 'test')}), flush=True)
    save_reports(output, results, top_ks, excel=options.all_models)
    failures = sum(result.get('status') == 'failed' for result in results)
    if failures:
        raise RuntimeError('{} checkpoints failed. See status/error columns in the saved reports.'.format(failures))
    return results


if __name__ == '__main__':
    main()
